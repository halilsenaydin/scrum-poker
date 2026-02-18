from datetime import datetime
from openpyxl import Workbook
from openpyxl.chart import BarChart, ScatterChart, Reference, Series
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule
import statistics
from django.contrib import admin
from django.http import HttpResponse
from django.utils.html import format_html
from django.utils.translation import gettext_lazy as _
from .models import Room, Participant, Task, TaskVote, Sprint
from .utils import BaseExcelExport
from poker_service.admin_site import admin_site

# ─────────────────────────────────────────────
# Excel adapter classes
# ─────────────────────────────────────────────
class RoomAdminExcel(BaseExcelExport):
    def get_sprints(self, queryset):
        room_ids = queryset.values_list('id', flat=True)

        return Sprint.objects.filter(
            room_id__in=room_ids
        ).select_related('room').order_by('room', '-created_at')

    def get_export_title(self, queryset):
        n = queryset.count()
        if n == 1:
            r = queryset.first()

            return f"SCRUM POKER ANALYTICS – {r.name} ({r.room_code})"
        
        return f"SCRUM POKER ANALYTICS – {n} Rooms"

    def _sheet_pivot_room_summary(self, wb, rooms):
        ws = wb.create_sheet("Room Summary")
        headers = [
            "Room Code", "Room Name", "Sprints", "Active Sprints",
            "Participants", "Active Participants", "Tasks", "Active Tasks",
            "Completed Tasks", "Total Votes", "Avg Consensus (%)",
            "Avg Vote Value", "High Risk Tasks", "Medium Risk Tasks",
            "Low Risk Tasks", "Total SP", "Avg SP per Task"
        ]

        ws.append(headers)
        self._style_header(ws, row=1)
        
        for r in rooms:
            active_participants = r.get_active_participants().count()
            all_participants = r.participants.count()
            
            sprints_count = r.sprints.count()
            active_sprints = r.sprints.filter(is_active=True).count()
            
            tasks = r.tasks.all()
            active_tasks = tasks.filter(is_active=True)
            completed_tasks = tasks.filter(status='completed')
            
            total_votes = sum(t.get_vote_count() for t in tasks)
            
            consensus_scores = []
            vote_values = []
            risk_distribution = {"HIGH": 0, "MEDIUM": 0, "LOW": 0}
            total_sp = 0
            
            for t in tasks:
                if t.sp:
                    total_sp += t.sp
                    
                if t.metrics and 'consensus_score' in t.metrics:
                    consensus_scores.append(t.metrics['consensus_score'])
                    
                    score = t.metrics['consensus_score']
                    if score >= 80:
                        risk_distribution["LOW"] += 1
                    elif score >= 50:
                        risk_distribution["MEDIUM"] += 1
                    else:
                        risk_distribution["HIGH"] += 1
                
                if t.metrics and 'average' in t.metrics:
                    vote_values.append(t.metrics['average'])
            
            avg_consensus = round(statistics.mean(consensus_scores), 2) if consensus_scores else 0
            avg_vote = round(statistics.mean(vote_values), 2) if vote_values else 0
            avg_sp_per_task = round(total_sp / tasks.count(), 2) if tasks.count() > 0 else 0
            
            ws.append([
                r.room_code,
                r.name,
                sprints_count,
                active_sprints,
                all_participants,
                active_participants,
                tasks.count(),
                active_tasks.count(),
                completed_tasks.count(),
                total_votes,
                avg_consensus,
                avg_vote,
                risk_distribution["HIGH"],
                risk_distribution["MEDIUM"],
                risk_distribution["LOW"],
                total_sp,
                avg_sp_per_task
            ])
        
        last_row = ws.max_row
        
        # Add conditional formatting
        if last_row > 1:
            ws.conditional_formatting.add(
                f'K2:K{last_row}',
                ColorScaleRule(
                    start_type='num', start_value=0, start_color=self.COLORS['danger'],
                    mid_type='num', mid_value=50, mid_color=self.COLORS['warning'],
                    end_type='num', end_value=100, end_color=self.COLORS['success']
                )
            )
            
            # Data bars for votes
            ws.conditional_formatting.add(
                f'J2:J{last_row}',
                DataBarRule(start_type='min', end_type='max', color=self.COLORS['info'])
            )
        
        self._auto_col_width(ws)
        
        # Add comparison charts
        if last_row > 2:
            self._add_room_comparison_charts(ws, last_row)
        
        self._freeze(ws)

    def _add_room_comparison_charts(self, ws, last_row):
        # Chart 1: Task Status Distribution
        chart1 = BarChart()
        chart1.type = "col"
        chart1.style = 10
        chart1.title = "Room Task Distribution"
        chart1.y_axis.title = 'Task Count'
        chart1.height = 12
        chart1.width = 20
        
        # Data: Total, Active, Completed
        data_ref = Reference(ws, min_col=7, min_row=1, max_col=9, max_row=min(last_row, 15))
        categories = Reference(ws, min_col=1, min_row=2, max_row=min(last_row, 15))
        
        chart1.add_data(data_ref, titles_from_data=True)
        chart1.set_categories(categories)
        chart1.grouping = "clustered"
        
        ws.add_chart(chart1, 'R2')
        
        # Chart 2: Consensus vs Votes scatter
        chart2 = ScatterChart()
        chart2.title = "Consensus vs Participation"
        chart2.style = 13
        chart2.x_axis.title = 'Total Votes'
        chart2.y_axis.title = 'Avg Consensus (%)'
        chart2.height = 12
        chart2.width = 20
        
        xvalues = Reference(ws, min_col=10, min_row=2, max_row=min(last_row, 15))
        yvalues = Reference(ws, min_col=11, min_row=2, max_row=min(last_row, 15))
        series = Series(yvalues, xvalues, title="Rooms")
        chart2.series.append(series)
        
        ws.add_chart(chart2, 'R20')

class SprintAdminExcel(BaseExcelExport):
    def get_sprints(self, queryset):
        return queryset.select_related('room').order_by('room', '-created_at')

    def get_export_title(self, queryset):
        n = queryset.count()
        if n == 1:
            s = queryset.first()
    
            return f"SPRINT ANALYTICS – {s.name} ({s.room.name})"
    
        return f"SPRINT ANALYTICS – {n} Sprints"

# ─────────────────────────────────────────────
# Shared export helper
# ─────────────────────────────────────────────
def _build_response(wb, filename):
    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument"
            ".spreadsheetml.sheet"
        )
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)

    return response

# ─────────────────────────────────────────────
# RoomAdmin
# ─────────────────────────────────────────────
class RoomAdmin(admin.ModelAdmin, RoomAdminExcel):
    list_display = [
        'room_code_badge',
        'name_display',
        'sprint_count',
        'participant_count',
        'task_count',
        'vote_count',
        'avg_consensus_display',
        'activity_indicator',
        'is_active',
        'created_at_display',
    ]
    list_filter   = ['is_active', 'created_at']
    search_fields = ['name', 'room_code']
    readonly_fields = ['room_code', 'created_at']
    fieldsets = (
        (_('admin_room_fieldset_base_info_label'), {
            'fields': ('name', 'room_code', 'room_password'),
        }),
        (_('admin_room_fieldset_status_label'), {
            'fields': ('is_active',),
        }),
        (_('admin_room_fieldset_time_info_label'), {
            'fields': ('created_at',),
            'classes': ('collapse',),
        }),
    )
    actions = [
        'activate_rooms',
        'deactivate_rooms',
        'export_basic_excel',
        'export_advanced_analytics',
        'export_full_analytics_with_charts',
    ]
    list_per_page = 25

    # ── Display Columns ──
    @admin.display(
        description=Room._meta.get_field("room_code").verbose_name,
        ordering="room_code",
    )
    def room_code_badge(self, obj):
        return format_html(
            '<span style="font-family:monospace;font-weight:bold;color:#366092;'
            'background:#f0f0f0;padding:3px 8px;border-radius:4px;">{}</span>',
            obj.room_code,
        )

    @admin.display(
        description=Room._meta.get_field("name").verbose_name,
        ordering="name",
    )
    def name_display(self, obj):
        e = "🚀" if obj.is_active else "💤"
        return format_html('{} {}', e, obj.name)

    def sprint_count(self, obj):
        total  = obj.sprints.count()
        active = obj.sprints.filter(is_active=True).count()
        if active:
            return format_html(
                '<span style="color:#22c55e;font-weight:bold;">🎯 {}</span> '
                '<span style="color:#666;font-size:11px;">({} {})</span>',
                total, active, _('admin_room_sprint_count_active_text')
            )
        return format_html('<span style="color:gray;">🎯 {}</span>', total)
    sprint_count.short_description = _('admin_room_sprint_count_label')

    def participant_count(self, obj):
        active = obj.get_participant_count()
        total  = obj.participants.count()
        if active:
            clr = '#22c55e' if active >= 5 else '#f59e0b' if active >= 3 else '#3b82f6'
            return format_html(
                '<span style="color:{};font-weight:bold;">👥 {}</span> '
                '<span style="color:#999;font-size:10px;">/ {}</span>',
                clr, active, total,
            )
        return format_html('<span style="color:gray;">0</span>')
    participant_count.short_description = _('admin_room_participants_label')

    def task_count(self, obj):
        total     = obj.tasks.count()
        completed = obj.tasks.filter(status='completed').count()
        if total:
            pct = round(completed / total * 100)
            clr = '#22c55e' if pct >= 70 else '#f59e0b' if pct >= 30 else '#3b82f6'
            return format_html(
                '<span style="color:{};font-weight:bold;">📋 {}</span> '
                '<span style="color:#666;font-size:11px;">({}/{} ✓ {}%)</span>',
                clr, total, completed, total, pct,
            )
        return format_html('<span style="color:gray;">0</span>')
    task_count.short_description = _('admin_room_tasks_label')

    def vote_count(self, obj):
        count = sum(t.get_vote_count() for t in obj.tasks.filter(is_active=True))
        if count:
            bar = min(count, 100)
            return format_html(
                '<div style="display:flex;align-items:center;gap:8px;">'
                '<span style="font-weight:bold;color:#8b5cf6;min-width:30px;">✓ {}</span>'
                '<div style="background:#e0e7ff;width:100px;height:12px;border-radius:6px;overflow:hidden;">'
                '<div style="background:#8b5cf6;width:{}%;height:100%;"></div>'
                '</div></div>',
                count, bar,
            )
        return format_html('<span style="color:gray;">0</span>')
    vote_count.short_description = _('admin_room_votes_label')

    def avg_consensus_display(self, obj):
        scores = [
            t.metrics['consensus_score']
            for t in obj.tasks.filter(is_active=True)
            if t.metrics and 'consensus_score' in t.metrics
        ]
        if scores:
            avg = round(statistics.mean(scores), 1)
            if avg >= 80:
                clr, em, lbl = '#22c55e', '✅', 'High'
            elif avg >= 50:
                clr, em, lbl = '#f59e0b', '⚠️',  'Med'
            else:
                clr, em, lbl = '#ef4444', '❌', 'Low'
            return format_html(
                '<span style="color:{};font-weight:bold;">{} {}% '
                '<span style="font-size:10px;opacity:.7;">({})</span></span>',
                clr, em, avg, lbl,
            )
        return format_html('<span style="color:gray;">–</span>')
    avg_consensus_display.short_description = _('admin_room_consensus_label')

    def activity_indicator(self, obj):
        recent = sum(
            t.votes.filter(voted_at__date=datetime.now().date()).count()
            for t in obj.tasks.all()
        )
        if recent >= 20:
            icon, tip = '🔥', _('admin_room_consensus_very_active_text')
        elif recent >= 10:
            icon, tip = '⚡', _('admin_room_consensus_active_text')
        elif recent >= 5:
            icon, tip = '📊', _('admin_room_consensus_moderate_text')
        else:
            icon, tip = '💤', _('admin_room_consensus_low_text')
        return format_html(
            '<span style="font-size:20px;" title="{}">{}</span>', tip, icon
        )
    activity_indicator.short_description = _('admin_room_activity_label')

    @admin.display(
        description=Room._meta.get_field("created_at").verbose_name,
        ordering="created_at",
    )
    def created_at_display(self, obj):
        if obj.created_at:
            return format_html(
                '<span style="color:#666;font-size:11px;">📅 {}</span>',
                obj.created_at.strftime('%Y-%m-%d %H:%M'),
            )
        return '–'

    # ── Actions ──
    def activate_rooms(self, request, queryset):
        n = queryset.update(is_active=True)
        self.message_user(request, f'{n} room(s) activated.', level='SUCCESS')
    activate_rooms.short_description = _('admin_room_action_activate_rooms_label')

    def deactivate_rooms(self, request, queryset):
        n = queryset.update(is_active=False)

        self.message_user(request, f'{n} room(s) deactivated.', level='WARNING')
    deactivate_rooms.short_description = _('admin_room_action_deactivate_rooms_label')

    def export_basic_excel(self, request, queryset):
        wb = Workbook()

        self._sheet_pivot_room_summary(wb, queryset)

        wb = self.export_basic(wb, queryset)
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        n_s = self.get_sprints(queryset).count()

        self.message_user(
            request,
            f'Basic export: {queryset.count()} room(s), {n_s} sprint(s).',
            level='SUCCESS',
        )

        return _build_response(wb, f'rooms_basic_{ts}.xlsx')
    export_basic_excel.short_description = _('admin_room_action_export_basic_excel_label')

    def export_advanced_analytics(self, request, queryset):
        wb = Workbook()

        self._sheet_pivot_room_summary(wb, queryset)

        wb = self.export_advanced(wb, queryset)
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        n_s = self.get_sprints(queryset).count()

        self.message_user(
            request,
            f'Analytics export: {queryset.count()} room(s), {n_s} sprint(s).',
            level='SUCCESS',
        )

        return _build_response(wb, f'rooms_analytics_{ts}.xlsx')
    export_advanced_analytics.short_description = _('admin_room_action_export_advanced_analytics_label')

    def export_full_analytics_with_charts(self, request, queryset):
        wb = Workbook()

        self._sheet_pivot_room_summary(wb, queryset)

        wb = self.export_full(wb, queryset)
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        n_s = self.get_sprints(queryset).count()
    
        self.message_user(
            request,
            f'Full export: {queryset.count()} room(s), {n_s} sprint(s).',
            level='SUCCESS',
        )
    
        return _build_response(wb, f'rooms_full_{ts}.xlsx')
    export_full_analytics_with_charts.short_description = _('admin_room_action_export_full_analytics_with_charts_label')

    def get_readonly_fields(self, request, obj=None):
        if obj:
            return ['room_code', 'created_at', 'room_password']
        return ['room_code', 'created_at']

admin_site.register(Room, RoomAdmin)

# ─────────────────────────────────────────────
# SprintAdmin
# ─────────────────────────────────────────────
class SprintAdmin(admin.ModelAdmin, SprintAdminExcel):
    list_display = [
        'sprint_name_display',
        'room_display',
        'is_active',
        'revealed_display',
        'task_count_display',
        'completion_display',
        'sp_display',
        'consensus_display',
        'created_at_display',
    ]
    list_filter   = ['is_active', 'revealed', 'room', 'created_at']
    search_fields = ['name', 'room__name', 'room__room_code']
    readonly_fields = ['created_at']
    fieldsets = (
        (_('admin_sprint_fieldset_base_info_label'), {
            'fields': ('room', 'name', 'tasks'),
        }),
        (_('admin_sprint_fieldset_status_label'), {
            'fields': ('is_active', 'revealed'),
        }),
        (_('admin_sprint_fieldset_time_info_label'), {
            'fields': ('created_at',),
            'classes': ('collapse',),
        }),
    )
    actions = [
        'activate_sprints',
        'deactivate_sprints',
        'reveal_all_votes',
        'reset_all_votes',
        'export_basic_excel',
        'export_advanced_analytics',
        'export_full_analytics_with_charts',
    ]
    list_per_page = 25
    filter_horizontal = ['tasks']

    # ── Display Columns ──
    @admin.display(
        description=Sprint._meta.get_field("name").verbose_name,
        ordering="name",
    )
    def sprint_name_display(self, obj):
        return format_html(
            '<span style="font-family:monospace;font-weight:bold;color:#366092;'
            'background:#f0f0f0;padding:3px 8px;border-radius:4px;">{}</span>',
            obj.name,
        )

    @admin.display(
        description=Room._meta.get_field("name").verbose_name,
        ordering="room",
    )
    def room_display(self, obj):
        return format_html(
            '<span style="font-family:monospace;color:#666;">{}</span> {}',
            obj.room.room_code, obj.room.name,
        )

    @admin.display(
        description=Sprint._meta.get_field("revealed").verbose_name
    )
    def revealed_display(self, obj):
        return format_html(
            '<span style="font-size:16px;">{}</span>',
            '👁️' if obj.revealed else '🔒',
        )

    @admin.display(
        description=Sprint._meta.get_field("tasks").verbose_name
    )
    def task_count_display(self, obj):
        total     = obj.tasks.count()
        completed = obj.tasks.filter(status='completed').count()
        if total:
            return format_html(
                '<span style="font-weight:bold;">📋 {}</span> '
                '<span style="color:#666;font-size:11px;">({} ✓)</span>',
                total, completed,
            )
        return format_html('<span style="color:gray;">0</span>')

    def completion_display(self, obj):
        total = obj.tasks.count()
        completed = obj.tasks.filter(status='completed').count()

        if not total:
            return '–'
    
        pct = round(completed / total * 100)
        clr = '#22c55e' if pct >= 70 else '#f59e0b' if pct >= 30 else '#ef4444'

        return format_html(
            '<div style="display:flex;align-items:center;gap:6px;">'
            '<span style="color:{};font-weight:bold;min-width:35px;">{}%</span>'
            '<div style="background:#e5e7eb;width:60px;height:8px;border-radius:4px;overflow:hidden;">'
            '<div style="background:{};width:{}%;height:100%;"></div>'
            '</div></div>',
            clr, pct, clr, pct,
        )
    completion_display.short_description = _('admin_sprint_completion_label')

    def sp_display(self, obj):
        tasks    = obj.tasks.all()
        total_sp = sum(t.sp for t in tasks if t.sp) or 0
        done_sp  = sum(t.sp for t in tasks.filter(status='completed') if t.sp) or 0

        if total_sp:
            return format_html(
                '<span style="color:#8b5cf6;font-weight:bold;">{}</span> '
                '<span style="color:#999;font-size:11px;">/ {}</span>',
                done_sp, total_sp,
            )

        return '–'
    sp_display.short_description = _('admin_sprint_sp_label')

    def consensus_display(self, obj):
        scores = [
            t.metrics['consensus_score']
            for t in obj.tasks.all()
            if t.metrics and 'consensus_score' in t.metrics
        ]

        if scores:
            avg = round(statistics.mean(scores), 1)

            if avg >= 80:
                return format_html('<span style="color:#22c55e;font-weight:bold;">✅ {}%</span>', avg)
            elif avg >= 50:
                return format_html('<span style="color:#f59e0b;font-weight:bold;">⚠️ {}%</span>', avg)
            return format_html('<span style="color:#ef4444;font-weight:bold;">❌ {}%</span>', avg)

        return '–'
    consensus_display.short_description = _('admin_sprint_consensus_label')

    @admin.display(
        description=Sprint._meta.get_field("created_at").verbose_name,
        ordering="created_at",
    )
    def created_at_display(self, obj):
        return obj.created_at.strftime('%Y-%m-%d') if obj.created_at else '–'

    # ── Actions ──
    def activate_sprints(self, request, queryset):
        n = queryset.update(is_active=True)
        self.message_user(request, f'{n} sprint(s) activated.', level='SUCCESS')
    activate_sprints.short_description = _('admin_sprint_action_activate_sprints_label')

    def deactivate_sprints(self, request, queryset):
        n = queryset.update(is_active=False)
        self.message_user(request, f'{n} sprint(s) deactivated.', level='WARNING')
    deactivate_sprints.short_description = _('admin_sprint_action_deactivate_sprints_label')

    def reveal_all_votes(self, request, queryset):
        for sprint in queryset:
            sprint.reveal_votes()
    
        self.message_user(
            request,
            f'Votes revealed for {queryset.count()} sprint(s).',
            level='SUCCESS',
        )
    reveal_all_votes.short_description = _('admin_sprint_action_reveal_all_votes_label')

    def reset_all_votes(self, request, queryset):
        for sprint in queryset:
            sprint.reset_votes()
    
        self.message_user(
            request,
            f'Votes reset for {queryset.count()} sprint(s).',
            level='WARNING',
        )
    reset_all_votes.short_description = _('admin_sprint_action_reset_all_votes_label')

    def export_basic_excel(self, request, queryset):
        wb = Workbook()
        wb = self.export_basic(wb, queryset)
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        tc = sum(s.tasks.count() for s in queryset)
    
        self.message_user(
            request,
            f'Basic export: {queryset.count()} sprint(s), {tc} task(s).',
            level='SUCCESS',
        )
    
        return _build_response(wb, f'sprints_basic_{ts}.xlsx')
    export_basic_excel.short_description = _('admin_sprint_action_export_basic_excel_label')

    def export_advanced_analytics(self, request, queryset):
        wb = Workbook()
        wb = self.export_advanced(wb, queryset)
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        tc = sum(s.tasks.count() for s in queryset)

        self.message_user(
            request,
            f'Analytics export: {queryset.count()} sprint(s), {tc} task(s).',
            level='SUCCESS',
        )

        return _build_response(wb, f'sprints_analytics_{ts}.xlsx')
    export_advanced_analytics.short_description = _('admin_sprint_action_export_advanced_analytics_label')

    def export_full_analytics_with_charts(self, request, queryset):
        wb = Workbook()
        wb = self.export_full(wb, queryset)
        ts = datetime.now().strftime('%Y%m%d_%H%M%S')
        tc = sum(s.tasks.count() for s in queryset)

        self.message_user(
            request,
            f'Full export: {queryset.count()} sprint(s), {tc} task(s).',
            level='SUCCESS',
        )

        return _build_response(wb, f'sprints_full_{ts}.xlsx')
    export_full_analytics_with_charts.short_description = _('admin_sprint_action_export_full_analytics_with_charts_label')

admin_site.register(Sprint, SprintAdmin)

class ParticipantAdmin(admin.ModelAdmin):
    list_display = ['user', 'room_display', 'vote_count', 'is_active']
    list_filter = ['is_active', 'room']
    search_fields = ['user__username', 'room__room_code']
    
    fieldsets = (
        (_('admin_participant_fieldset_base_info_label'), {
            'fields': ('room', 'user')
        }),
    )

    @admin.display(
        description=Participant._meta.get_field("room").verbose_name,
        ordering="room",
    )
    def room_display(self, obj):
        return format_html(
            '<span style="font-family:monospace;color:#666;">{}</span> {}',
            obj.room.room_code, obj.room.name,
        )

    def room_code(self, obj):
        return obj.room.room_code
    room_code.short_description = _('admin_participant_room_code_label')
    room_code.admin_order_field = 'room__room_code'
    
    def vote_count(self, obj):
        count = obj.votes.count()
        if count > 0:
            return format_html('<span style="color: green; font-weight: bold;">{}</span>', count)
        return format_html('<span style="color: gray;">0</span>')
    vote_count.short_description = _('admin_participant_vote_count_label')

admin_site.register(Participant, ParticipantAdmin)

class TaskAdmin(admin.ModelAdmin):
    list_display = ['title', 'room_code', 'sp', 'status', 'vote_count', 'consensus_indicator', 'is_active', 'created_at']
    list_filter = ['status', 'is_active', 'created_at', 'room']
    search_fields = ['title', 'description', 'room__room_code']
    readonly_fields = ['metrics', 'created_at', 'updated_at', 'metrics_display']
    
    fieldsets = (
        (_('admin_task_fieldset_base_info_label'), {
            'fields': ('room', 'title', 'description')
        }),
        (_('admin_task_fieldset_estimate_info_label'), {
            'fields': ('sp', 'status', 'metrics_display')
        }),
        (_('admin_task_fieldset_time_info_label'), {
            'fields': ('created_at', 'updated_at'),
            'classes': ('collapse',)
        }),
        (_('admin_task_fieldset_technical_info_label'), {
            'fields': ('metrics',),
            'classes': ('collapse',)
        }),
    )
    
    def room_code(self, obj):
        return obj.room.room_code
    room_code.short_description = _('admin_task_room_code_label')
    room_code.admin_order_field = 'room__room_code'
    
    def vote_count(self, obj):
        count = obj.get_vote_count()
        total = obj.room.get_participant_count()
        
        if total > 0:
            percentage = (count / total) * 100
            color = 'green' if percentage == 100 else 'orange' if percentage >= 50 else 'red'
            return format_html(
                '<span style="color: {}; font-weight: bold;">{}/{} ({}%)</span>',
                color, count, total, int(percentage)
            )
        return format_html('<span style="color: gray;">{}/0</span>', count)
    vote_count.short_description = _('admin_task_vote_count_label')
    
    def consensus_indicator(self, obj):
        if not obj.metrics or 'consensus_score' not in obj.metrics:
            return format_html('<span style="color: gray;">-</span>')
        
        score = obj.metrics['consensus_score']
        if score >= 80:
            color = 'green'
            icon = '✓'
        elif score >= 50:
            color = 'orange'
            icon = '~'
        else:
            color = 'red'
            icon = '✗'
        
        return format_html(
            '<span style="color: {}; font-weight: bold;">{} {}%</span>',
            color, icon, int(score)
        )
    consensus_indicator.short_description = _('admin_task_consensus_label')
    
    def metrics_display(self, obj):
        if not obj.metrics:
            return format_html('<span style="color: gray;">{}</span>', _("admin_task_metrics_no_votes"))
        
        html = '<table style="border-collapse: collapse;">'
        
        metrics_map = {
            'total_votes': _('admin_task_metrics_total_votes'),
            'average': _('admin_task_metrics_average'),
            'median': _('admin_task_metrics_median'),
            'min': _('admin_task_metrics_min'),
            'max': _('admin_task_metrics_max'),
            'std_dev': _('admin_task_metrics_std_dev'),
            'consensus_score': _('admin_task_metrics_consensus_score'),
        }
        
        for key, label in metrics_map.items():
            if key in obj.metrics:
                value = obj.metrics[key]
                html += f'<tr><td style="padding: 4px 8px; font-weight: bold;">{label}:</td><td style="padding: 4px 8px;">{value}</td></tr>'
        
        if 'distribution' in obj.metrics:
            html += f'<tr><td style="padding: 4px 8px; font-weight: bold;">{_("admin_task_metrics_distribution")}:</td><td style="padding: 4px 8px;">'
            for vote_val, count in sorted(obj.metrics['distribution'].items()):
                html += f'{vote_val}: {count} {_("admin_task_metrics_vote")}<br>'
            html += '</td></tr>'
        
        html += '</table>'
        return format_html(html)
    metrics_display.short_description = _('admin_task_metrics_display_label')
    
    actions = ['recalculate_metrics', 'reset_task_votes']
    
    def recalculate_metrics(self, request, queryset):
        for task in queryset:
            task.calculate_metrics()
            task.save()
        self.message_user(request, f"{queryset.count()} {_('admin_task_recalculated_metrics')}")
    recalculate_metrics.short_description = _('admin_task_action_recalculate_metrics')
    
    def reset_task_votes(self, request, queryset):
        for task in queryset:
            task.reset_votes()
        self.message_user(request, f"{queryset.count()} {_('admin_task_reset_task_votes')}")
    reset_task_votes.short_description = _('admin_task_action_reset_task_votes')

admin_site.register(Task, TaskAdmin)

class TaskVoteAdmin(admin.ModelAdmin):
    list_display = ['participant_name', 'task_title', 'vote', 'is_numeric_vote', 'voted_at']
    list_filter = ['voted_at', 'vote', 'task__room']
    search_fields = ['participant__name', 'task__title', 'vote']
    readonly_fields = ['voted_at']
    fieldsets = (
        (_('admin_task_vote_fieldset_vote_info_label'), {
            'fields': ('task', 'participant', 'vote')
        }),
        (_('admin_task_vote_fieldset_time_info_label'), {
            'fields': ('voted_at',),
            'classes': ('collapse',)
        }),
    )
    
    def participant_name(self, obj):
        return str(obj.participant.user)

    participant_name.short_description = _('admin_task_vote_participant_name_label')
    participant_name.admin_order_field = 'participant__name'
    
    def task_title(self, obj):
        return obj.task.title[:50] + '...' if len(obj.task.title) > 50 else obj.task.title
    
    task_title.short_description = _('admin_task_vote_task_title_label')
    task_title.admin_order_field = 'task__title'
    
    def is_numeric_vote(self, obj):
        if obj.is_numeric():
            return format_html('<span style="color: green;">✓ {}</span>', _('admin_task_vote_numeric_vote'))
        return format_html('<span style="color: orange;">{} ({0})</span>', _('admin_task_vote_custom_vote'), obj.vote)
    is_numeric_vote.short_description = _('admin_task_vote_is_numeric_vote_label')

admin_site.register(TaskVote, TaskVoteAdmin)
